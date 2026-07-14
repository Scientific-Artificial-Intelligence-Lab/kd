
from __future__ import annotations

import importlib
import math
import re
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pytest
import torch

from kd.data.schema import PDEDataset


def _under_report_dimension(path: Path) -> None:
    scratch = path.with_suffix(".tmp.xlsx")
    with (
        zipfile.ZipFile(path) as zin,
        zipfile.ZipFile(scratch, "w", zipfile.ZIP_DEFLATED) as zout,
    ):
        for item in zin.namelist():
            data = zin.read(item)
            if item.endswith("sheet1.xml"):
                text = data.decode("utf8")
                patched = re.sub(r'<dimension ref="[^"]*"/>', '<dimension ref="A1"/>', text)
                if patched == text:
                    patched = re.sub(
                        r"(<worksheet[^>]*>)", r'\1<dimension ref="A1"/>', text, count=1
                    )
                data = patched.encode("utf8")
            zout.writestr(item, data)
    scratch.replace(path)


def _write_workbook(
    path: Path,
    rows: list[tuple[object, ...]],
    *,
    sheet_name: str = "samples",
) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    workbook.close()


def _read_xlsx_columns(
    path: Path,
    column_names: tuple[str, ...],
    **kwargs: object,
) -> dict[str, list[float]]:
    module = importlib.import_module("kd.data.xlsx")
    reader: Callable[..., dict[str, list[float]]] = module.read_xlsx_columns
    return reader(path, column_names, **kwargs)


def _from_xlsx(path: Path, **kwargs: Any) -> PDEDataset:
    factory: Callable[..., PDEDataset] = PDEDataset.from_xlsx
    return factory(path, **kwargs)


@pytest.fixture
def sample_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "scatter.xlsx"
    _write_workbook(
        path,
        [
            ("x", "y", "values", "unused"),
            (0.0, 10.0, 1.5, "ignored"),
            (1.0, 11.0, "Indeterminate", "ignored"),
            (2.0, 12.0, None, "ignored"),
            (3.0, 13.0, 4.5, "ignored"),
        ],
    )
    return path


class TestReadXlsxColumns:

    @pytest.mark.unit
    def test_header_name_mapping_reads_selected_columns_in_request_order(
        self, sample_workbook: Path
    ) -> None:
        columns = _read_xlsx_columns(sample_workbook, ("values", "x"))

        assert list(columns) == ["values", "x"]
        assert columns["x"] == [0.0, 1.0, 2.0, 3.0]
        assert columns["values"][0] == 1.5
        assert math.isnan(columns["values"][1])
        assert math.isnan(columns["values"][2])
        assert columns["values"][3] == 4.5

    @pytest.mark.unit
    def test_missing_requested_header_lists_available_headers(
        self, sample_workbook: Path
    ) -> None:
        with pytest.raises(ValueError, match=r"available headers.*x.*y.*values"):
            _read_xlsx_columns(sample_workbook, ("not-a-column",))

    @pytest.mark.unit
    def test_duplicate_requested_header_names_ambiguous_column(
        self, tmp_path: Path
    ) -> None:
        path = tmp_path / "duplicate-header.xlsx"
        _write_workbook(
            path,
            [("x", "values", "values"), (0.0, 1.0, 2.0)],
        )

        with pytest.raises(ValueError, match="values"):
            _read_xlsx_columns(path, ("values",))

    @pytest.mark.unit
    def test_unexpected_non_numeric_value_names_row_column_and_value(
        self, tmp_path: Path
    ) -> None:
        path = tmp_path / "bad-value.xlsx"
        _write_workbook(path, [("x", "values"), (0.0, "not-a-number")])

        with pytest.raises(
            ValueError,
            match=r"row.*2.*values.*not-a-number",
        ):
            _read_xlsx_columns(path, ("x", "values"))

    @pytest.mark.unit
    def test_sheet_name_and_zero_based_sheet_position_select_requested_sheet(
        self, tmp_path: Path
    ) -> None:
        path = tmp_path / "sheets.xlsx"
        _write_workbook(path, [("x",), (1.0,)], sheet_name="first")
        workbook = openpyxl.load_workbook(path)
        second = workbook.create_sheet("second")
        second.append(("x",))
        second.append((2.0,))
        workbook.save(path)
        workbook.close()

        assert _read_xlsx_columns(path, ("x",), sheet="second") == {"x": [2.0]}
        assert _read_xlsx_columns(path, ("x",), sheet=0) == {"x": [1.0]}

    @pytest.mark.unit
    def test_under_reported_dimension_tag_still_reads_all_columns(
        self, tmp_path: Path
    ) -> None:
        path = tmp_path / "lazy-dimension.xlsx"
        _write_workbook(path, [("x", "y", "values"), (0.0, 10.0, 1.5), (1.0, 11.0, 2.5)])
        _under_report_dimension(path)

        columns = _read_xlsx_columns(path, ("x", "y", "values"))

        assert columns == {"x": [0.0, 1.0], "y": [10.0, 11.0], "values": [1.5, 2.5]}

    @pytest.mark.unit
    def test_boolean_cell_is_rejected_as_malformed(self, tmp_path: Path) -> None:
        path = tmp_path / "bool-cell.xlsx"
        _write_workbook(path, [("x", "values"), (0.0, True)])

        with pytest.raises(ValueError, match=r"boolean cell at row.*2.*values"):
            _read_xlsx_columns(path, ("x", "values"))


class TestPDEDatasetFromXlsx:

    @pytest.mark.unit
    def test_drop_na_discards_each_row_with_na_in_any_selected_column(
        self, sample_workbook: Path
    ) -> None:
        dataset = _from_xlsx(
            sample_workbook,
            coords={"x": "x", "y": "y"},
            fields={"u": "values"},
            lhs="",
            name="steady-samples",
            ground_truth="u_xx + u_yy = 0",
        )

        assert dataset.name == "steady-samples"
        assert dataset.lhs_field == ""
        assert dataset.lhs_axis == ""
        assert dataset.lhs_order == 0
        assert dataset.ground_truth == "u_xx + u_yy = 0"
        assert dataset.get_shape() == (2,)
        torch.testing.assert_close(
            dataset.get_coords("x"), torch.tensor([0.0, 3.0], dtype=torch.float64)
        )
        torch.testing.assert_close(
            dataset.get_field("u"), torch.tensor([1.5, 4.5], dtype=torch.float64)
        )

    @pytest.mark.unit
    def test_nonempty_lhs_uses_existing_lhs_parser(self, sample_workbook: Path) -> None:
        dataset = _from_xlsx(
            sample_workbook,
            coords={"x": "x", "t": "y"},
            fields={"u": "values"},
            lhs="u_t",
        )

        assert dataset.lhs_field == "u"
        assert dataset.lhs_axis == "t"
        assert dataset.lhs_order == 1

    @pytest.mark.unit
    def test_drop_na_false_preserves_nan_rows(self, sample_workbook: Path) -> None:
        dataset = _from_xlsx(
            sample_workbook,
            coords={"x": "x", "y": "y"},
            fields={"u": "values"},
            lhs="",
            drop_na=False,
        )

        assert dataset.get_shape() == (4,)
        values = dataset.get_field("u")
        assert np.isnan(values[1].item())
        assert np.isnan(values[2].item())
