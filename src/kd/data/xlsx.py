
from __future__ import annotations

import warnings
from collections.abc import Iterator, Sequence
from contextlib import contextmanager
from numbers import Real
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook

from kd.data._columns import column_indices

_NA_VALUES: tuple[str, ...] = ("Indeterminate",)


@contextmanager
def _worksheet(path: str | Path, sheet: str | int | None) -> Iterator[Any]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        worksheet = _select_sheet(workbook, sheet)





        if hasattr(worksheet, "reset_dimensions"):
            worksheet.reset_dimensions()
        yield worksheet
    finally:
        workbook.close()


def _header(worksheet: Any, header_row: int, path: str | Path) -> tuple[object, ...]:
    header = next(
        worksheet.iter_rows(
            min_row=header_row + 1,
            max_row=header_row + 1,
            values_only=True,
        ),
        None,
    )
    if header is None:
        raise ValueError(f"XLSX file has no header row {header_row}: {path}")
    return tuple(header)


def read_xlsx_columns(
    path: str | Path,
    column_names: Sequence[str],
    *,
    sheet: str | int | None = None,
    header_row: int = 0,
    na_values: Sequence[str] = _NA_VALUES,
) -> dict[str, list[float]]:
    if header_row < 0:
        raise ValueError(f"header_row must be >= 0, got {header_row}")

    with _worksheet(path, sheet) as worksheet:
        header_to_index = column_indices(
            _header(worksheet, header_row, path), column_names, kind="XLSX"
        )

        columns: dict[str, list[float]] = {name: [] for name in column_names}
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 2, values_only=True),
            start=header_row + 2,
        ):
            for name, values in columns.items():
                index = header_to_index[name]
                if index >= len(row):
                    raise ValueError(
                        f"row {row_index} is shorter than the header: missing "
                        f"column {name!r} (index {index}, row width {len(row)})"
                    )
                values.append(_cell_as_float(row[index], row_index, name, na_values))
        return columns


def read_xlsx_inventory_columns(path: Path) -> dict[str, np.ndarray]:
    with _worksheet(path, None) as worksheet:
        header = _header(worksheet, 0, path)
        if not all(isinstance(name, str) and name for name in header):
            raise ValueError(f"XLSX headers must be non-empty text, got {header}")
        names = [str(name) for name in header]
        column_indices(names, names, kind="XLSX")
        cells: list[list[object]] = [[] for _ in names]
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            if any(value is not None for value in row[len(names):]):
                raise ValueError(f"XLSX row {row_index} has cells beyond the header")
            for index, column in enumerate(cells):
                column.append(row[index] if index < len(row) else None)
    return {
        name: _inventory_column(column, name)
        for name, column in zip(names, cells, strict=True)
    }


def _inventory_column(cells: list[object], name: str) -> np.ndarray:
    try:
        return np.asarray(
            [
                _cell_as_float(value, row, name, _NA_VALUES)
                for row, value in enumerate(cells, start=2)
            ],
            dtype=np.float64,
        )
    except ValueError:
        return np.asarray(["" if value is None else str(value) for value in cells])


def _select_sheet(workbook: Any, sheet: str | int | None) -> Any:
    if sheet is None:
        return workbook.active
    if isinstance(sheet, str):
        return workbook[sheet]
    if isinstance(sheet, int):
        return workbook.worksheets[sheet]
    raise TypeError(f"sheet must be str, int, or None, got {type(sheet).__name__}")


def _cell_as_float(
    value: object,
    row_index: int,
    column_name: str,
    na_values: Sequence[str],
) -> float:
    if value is None or value in na_values:
        return float("nan")


    if isinstance(value, bool):
        raise ValueError(
            f"unexpected boolean cell at row {row_index}, column {column_name!r}: "
            f"{value!r}"
        )
    if isinstance(value, Real):
        return float(value)
    raise ValueError(
        f"unexpected non-numeric cell at row {row_index}, column {column_name!r}: "
        f"{value!r}"
    )
